"""Memory for an AI-led job hunt: the AI picks target companies and checks each one with its own web
search; this keeps the to-do list of companies and the verified jobs, and writes them to one Excel file.

State lives in hunt_state.json (git-ignored) so a long hunt can continue across chats.
Every saved job needs a link. The link is opened once; one that isn't found or can't be checked is
flagged in the Excel file and in the tool's reply (never hidden), since blocked sites also answer 404.
"""

import json
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font

import linkedin_bot as bot

STATE = Path(__file__).parent / "hunt_state.json"
EXCEL = bot.OUTPUT_DIR / "job_hunt.xlsx"
COLUMNS = ["Fit", "Title", "Company", "Location", "Type", "Term", "Deadline", "Eligibility", "Why it fits",
           "Requirements", "Link", "Link check", "Found"]
TIERS = ("reach", "target", "safe")
UA = {"User-Agent": "job-alert-bot/1.0 (personal job search; link check)"}


def _load():
    try:
        return json.loads(STATE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {"companies": [], "jobs": []}


def _save(state):
    STATE.write_text(json.dumps(state, indent=1), encoding="utf-8")


def _key(name):
    return re.sub(r"\W+", " ", name.lower()).strip()


def queue_companies(companies):
    """companies: [{name, tier, why}]. Adds new ones; returns how many were new."""
    state = _load()
    have = {_key(c["name"]) for c in state["companies"]}
    added = 0
    for c in companies:
        name = (c.get("name") or "").strip()
        if not name or _key(name) in have:
            continue
        tier = c.get("tier", "target")
        state["companies"].append({"name": name, "tier": tier if tier in TIERS else "target",
                                   "why": c.get("why", ""), "status": "pending", "note": ""})
        have.add(_key(name))
        added += 1
    _save(state)
    return added


def status(batch=8):
    state = _load()
    order = {t: i for i, t in enumerate(("target", "reach", "safe"))}
    pending = sorted((c for c in state["companies"] if c["status"] == "pending"), key=lambda c: order[c["tier"]])
    return {"companies_total": len(state["companies"]), "checked": len(state["companies"]) - len(pending),
            "pending": len(pending), "next_batch": [{k: c[k] for k in ("name", "tier", "why")} for c in pending[:batch]],
            "jobs_saved": len(state["jobs"]), "excel": str(EXCEL)}


def _check_link(url):
    try:
        r = requests.get(url, headers=UA, timeout=10, allow_redirects=True, stream=True)
        r.close()
    except requests.RequestException:
        return "Couldn't open it, check by hand"
    if r.status_code in (404, 410):
        return "NOT FOUND (404): may be a dead or invented link, or a site that blocks checks. Verify by hand"
    return "OK" if r.status_code < 400 else "Site blocks automatic checks. Open it by hand"


def save_findings(jobs, checked=()):
    """jobs: [{title, company, link, location, type, term, deadline, eligibility, why, requirements, fit}].
    checked: [{name, note}] companies that were looked at (with or without jobs). Returns a summary."""
    state = _load()
    rejected, ok = [], []
    for j in jobs:
        link = (j.get("link") or "").strip()
        if not (j.get("title") and j.get("company")) or not re.match(r"https?://", link):
            rejected.append({"title": j.get("title", "?"), "reason": "needs title, company and an http(s) link of the real posting"})
        else:
            ok.append(j)
    with ThreadPoolExecutor(6) as pool:
        checks = list(pool.map(lambda j: _check_link(j["link"].strip()), ok))
    have = {x["link"].split("#")[0] for x in state["jobs"]}
    saved, warnings = 0, []
    for j, check in zip(ok, checks):
        link = j["link"].strip()
        if check != "OK":
            warnings.append({"title": j["title"], "company": j["company"], "link": link, "link_check": check})
        if link.split("#")[0] not in have:
            state["jobs"].append({
                "fit": int(j.get("fit") or 0), "title": j["title"], "company": j["company"],
                "location": j.get("location", ""), "type": j.get("type", ""), "term": j.get("term", ""),
                "deadline": j.get("deadline", ""), "eligibility": j.get("eligibility", ""), "why": j.get("why", ""),
                "requirements": j.get("requirements", ""), "link": link, "link_check": check,
                "found": datetime.now().strftime("%Y-%m-%d")})
            have.add(link.split("#")[0])
            saved += 1
    done = {_key(c["name"]): c.get("note", "") for c in checked}
    for c in state["companies"]:
        if _key(c["name"]) in done:
            c["status"], c["note"] = "checked", done[_key(c["name"])]
    for j in state["jobs"]:       # a company with a saved job counts as checked
        for c in state["companies"]:
            if _key(c["name"]) == _key(j["company"]):
                c["status"] = "checked"
    state["jobs"].sort(key=lambda x: x["fit"], reverse=True)
    _save(state)
    path = _write_excel(state["jobs"])
    return {"saved": saved, "rejected": rejected, "link_warnings": warnings, "total_jobs": len(state["jobs"]), "excel": path,
            "pending_companies": sum(c["status"] == "pending" for c in state["companies"])}


def _write_excel(jobs):
    bot.OUTPUT_DIR.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Job hunt"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for j in jobs:
        ws.append([j["fit"], j["title"], j["company"], j["location"], j["type"], j["term"], j["deadline"],
                   j["eligibility"], j["why"], j["requirements"], j["link"], j["link_check"], j["found"]])
    for row in ws.iter_rows(min_row=2, min_col=11, max_col=11):
        row[0].hyperlink = row[0].value
        row[0].font = Font(color="0563C1", underline="single")
    for row in ws.iter_rows(min_row=2, min_col=12, max_col=12):
        if row[0].value != "OK":
            row[0].font = Font(color="C00000", bold=True)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 55)
    ws.freeze_panes = "A2"
    try:
        wb.save(EXCEL)
        return str(EXCEL)
    except PermissionError:     # open in Excel: save beside it instead of losing the data
        alt = EXCEL.with_name(f"job_hunt_{datetime.now():%H%M%S}.xlsx")
        wb.save(alt)
        return str(alt)
